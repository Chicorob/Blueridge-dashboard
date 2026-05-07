"""
BlueRidge Life Sciences Dashboard - PowerPoint Export
Exports dashboard views using the official BRLS PowerPoint template.
"""

import io
import logging
import os
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import plotly.graph_objects as go
from config import (
    COMPANY_NAME, DIVISIONS, DIVISION_COLORS,
    METRICS, CURRENCY_METRICS,
)
from data_manager import format_value, compute_delta, determine_display_unit

logger = logging.getLogger(__name__)


class ChartRenderError(RuntimeError):
    """Raised when a Plotly figure cannot be rendered to an image (Kaleido missing or failed)."""


# Brand colors from template
BLUE = RGBColor(0x00, 0x66, 0xB3)
DARK_BLUE = RGBColor(0x00, 0x3D, 0x6B)
GREEN = RGBColor(0x7A, 0xB6, 0x48)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x2C, 0x3E, 0x50)
RED = RGBColor(0xE7, 0x4C, 0x3C)
POS_GREEN = RGBColor(0x27, 0xAE, 0x60)
SUBTLE_GRAY = RGBColor(0x56, 0x65, 0x73)

# Template layout indices
LAYOUT_TITLE_SLIDE = 0       # Branded title slide with background image
LAYOUT_TITLE_AND_CONTENT = 1 # Title bar + content area + footer
LAYOUT_TITLE_ONLY = 5        # Title bar + open area + footer
LAYOUT_BLANK = 6             # Footer/logo only, open canvas


def _get_template_path():
    """Get path to the BRLS PowerPoint template."""
    base = os.path.dirname(os.path.abspath(__file__))
    p = os.path.join(base, "assets", "BRLS_PPT_Template.pptx")
    if os.path.exists(p):
        return p
    return None


def create_presentation():
    """Create a new presentation from the BRLS template."""
    template_path = _get_template_path()
    if template_path:
        prs = Presentation(template_path)
        # Remove the placeholder "Template" slide
        if len(prs.slides) > 0:
            rId = prs.slides._sldIdLst[0].get(
                '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
            )
            prs.part.drop_rel(rId)
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    return prs


def _add_title_slide(prs, title_text):
    """Add a branded title slide using the template's Title Slide layout."""
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_SLIDE])

    # Set the title placeholder text
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 0:  # Title placeholder
            shape.text = title_text
            for para in shape.text_frame.paragraphs:
                para.font.size = Pt(40)
                para.font.bold = True
                para.font.name = "Times New Roman"
            break

    return slide


def _add_content_slide(prs, title_text, subtitle_text=""):
    """Add a slide using Title Only layout - title bar with open content area."""
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])

    # Set the title
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 0:  # Title
            shape.text = title_text
            for para in shape.text_frame.paragraphs:
                para.font.size = Pt(24)
                para.font.bold = True
                para.font.name = "Times New Roman"
                para.font.color.rgb = DARK_BLUE
            break

    # Add subtitle below the title bar
    if subtitle_text:
        stx = slide.shapes.add_textbox(Inches(1.5), Inches(1.15), Inches(10), Inches(0.35))
        sp = stx.text_frame.paragraphs[0]
        sp.text = subtitle_text
        sp.font.size = Pt(13)
        sp.font.color.rgb = SUBTLE_GRAY
        sp.font.name = "Times New Roman"

    return slide


def _add_blank_slide(prs):
    """Add a blank slide with just footer branding."""
    return prs.slides.add_slide(prs.slide_layouts[LAYOUT_BLANK])


def _add_kpi_card(slide, left, top, width, height, label, value, delta=None):
    """Add a KPI card shape to a slide."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF0, 0xF7, 0xFC)
    shape.line.color.rgb = RGBColor(0xAE, 0xD6, 0xF1)
    shape.line.width = Pt(1)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.1)
    tf.margin_left = Inches(0.15)

    # Label
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(11)
    p.font.color.rgb = SUBTLE_GRAY
    p.font.name = "Times New Roman"
    p.font.bold = False

    # Value
    p2 = tf.add_paragraph()
    p2.text = value
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = DARK_BLUE
    p2.font.name = "Times New Roman"

    # Delta
    if delta is not None:
        p3 = tf.add_paragraph()
        arrow = "▲" if delta >= 0 else "▼"
        p3.text = f"{arrow} {abs(delta):.1f}% vs prior year"
        p3.font.size = Pt(10)
        p3.font.color.rgb = POS_GREEN if delta >= 0 else RED
        p3.font.name = "Times New Roman"


def _chart_to_image_bytes(fig, width=900, height=450):
    """Convert a Plotly figure to PNG bytes.

    Raises ChartRenderError when the underlying Kaleido renderer is missing or
    fails. Callers can catch this to fall back to a chart-free export.
    """
    try:
        img_bytes = fig.to_image(format="png", width=width, height=height, scale=2)
        return io.BytesIO(img_bytes)
    except Exception as e:
        logger.warning("Kaleido scale=2 render failed (%s); retrying at scale=1", e)
        try:
            img_bytes = fig.to_image(format="png", width=width, height=height, scale=1)
            return io.BytesIO(img_bytes)
        except Exception as e2:
            logger.error("Chart render failed; install kaleido for PPTX chart export: %s", e2)
            raise ChartRenderError(str(e2)) from e2


def _make_bar_chart_fig(comparison_data, metric, unit=None):
    """Create a Plotly bar chart figure for export."""
    divisions = list(comparison_data.keys())
    values = list(comparison_data.values())
    colors = [DIVISION_COLORS.get(d, "#0066B3") for d in divisions]

    fig = go.Figure(data=[go.Bar(
        x=divisions, y=values,
        marker_color=colors,
        text=[format_value(v, metric, unit=unit) for v in values],
        textposition="outside",
        textfont=dict(size=16, family="Times New Roman"),
    )])
    fig.update_layout(
        font=dict(family="Times New Roman", size=15),
        plot_bgcolor="white",
        paper_bgcolor="white",
        margin=dict(l=60, r=40, t=20, b=60),
        yaxis=dict(gridcolor="#E8E8E8", title=""),
        xaxis=dict(title="", tickfont=dict(size=14)),
    )
    return fig


def _make_trend_chart_fig(trend_data_dict, metric):
    """Create a Plotly trend line chart figure for export."""
    fig = go.Figure()
    for entity, data in trend_data_dict.items():
        if data is not None and not data.empty:
            color = DIVISION_COLORS.get(entity, "#0066B3")
            fig.add_trace(go.Scatter(
                x=data["Date"], y=data[metric],
                name=entity, mode="lines+markers",
                line=dict(color=color, width=2.5),
                marker=dict(size=5),
            ))

    fig.update_layout(
        font=dict(family="Times New Roman", size=15),
        plot_bgcolor="white",
        paper_bgcolor="white",
        margin=dict(l=60, r=40, t=20, b=80),
        yaxis=dict(gridcolor="#E8E8E8", title=""),
        xaxis=dict(title=""),
        legend=dict(
            orientation="h", yanchor="bottom", y=-0.22,
            xanchor="center", x=0.5,
            font=dict(size=13, family="Times New Roman"),
        ),
    )
    return fig


# ─── Slide Export Functions ───

def export_executive_summary(prs, current_data, prior_data, time_range, entity, unit=None):
    """Export the executive summary / KPI overview slide."""
    subtitle = f"{entity}  |  {time_range}  |  Generated {datetime.now().strftime('%B %d, %Y')}"
    slide = _add_content_slide(prs, "Executive Summary", subtitle)

    metrics_list = list(METRICS.keys())

    # Determine unit if not provided
    if unit is None:
        currency_vals = [current_data.get(m, 0) for m in CURRENCY_METRICS]
        unit = determine_display_unit(currency_vals)

    # Layout: 5 cards per row
    cards_per_row = 5
    card_width = Inches(2.2)
    card_height = Inches(1.2)
    start_left = Inches(0.4)
    gap = Inches(0.15)

    for i, metric in enumerate(metrics_list):
        row = i // cards_per_row
        col = i % cards_per_row
        left = start_left + col * (card_width + gap)
        top = Inches(1.65) + row * (card_height + gap)

        val = current_data.get(metric, 0)
        prior_val = prior_data.get(metric, 0)
        delta = compute_delta(val, prior_val)
        if metric in CURRENCY_METRICS:
            formatted = format_value(val, metric, unit=unit)
        else:
            formatted = format_value(val, metric)

        _add_kpi_card(slide, left, top, card_width, card_height, metric, formatted, delta)

    return slide


def export_division_comparison(prs, comparison_data, metric, time_range, unit=None):
    """Export a division comparison bar chart slide."""
    slide = _add_content_slide(
        prs,
        f"Division Comparison: {metric}",
        f"{time_range}  |  All Divisions"
    )

    chart_unit = unit if metric in CURRENCY_METRICS else None
    fig = _make_bar_chart_fig(comparison_data, metric, unit=chart_unit)
    img_stream = _chart_to_image_bytes(fig, width=1100, height=480)
    if img_stream:
        slide.shapes.add_picture(
            img_stream, Inches(0.8), Inches(1.6),
            width=Inches(11.2), height=Inches(5.0)
        )

    return slide


def export_trend_chart(prs, trend_data_dict, metric, time_range):
    """Export a trend line chart slide for one or more entities."""
    slide = _add_content_slide(
        prs,
        f"Trend Analysis: {metric}",
        f"{time_range}"
    )

    fig = _make_trend_chart_fig(trend_data_dict, metric)
    img_stream = _chart_to_image_bytes(fig, width=1100, height=480)
    if img_stream:
        slide.shapes.add_picture(
            img_stream, Inches(0.8), Inches(1.6),
            width=Inches(11.2), height=Inches(5.0)
        )

    return slide


def export_division_detail(prs, division_data, prior_data, division, time_range, unit=None):
    """Export a detailed division slide with all KPIs."""
    subtitle = f"{division}  |  {time_range}"
    slide = _add_content_slide(prs, f"Division Detail: {division}", subtitle)

    metrics_list = list(METRICS.keys())

    if unit is None:
        currency_vals = [division_data.get(m, 0) for m in CURRENCY_METRICS]
        unit = determine_display_unit(currency_vals)

    cards_per_row = 5
    card_width = Inches(2.2)
    card_height = Inches(1.2)
    start_left = Inches(0.4)
    gap = Inches(0.15)

    for i, metric in enumerate(metrics_list):
        row = i // cards_per_row
        col = i % cards_per_row
        left = start_left + col * (card_width + gap)
        top = Inches(1.65) + row * (card_height + gap)

        val = division_data.get(metric, 0)
        prior_val = prior_data.get(metric, 0)
        delta = compute_delta(val, prior_val)
        if metric in CURRENCY_METRICS:
            formatted = format_value(val, metric, unit=unit)
        else:
            formatted = format_value(val, metric)

        _add_kpi_card(slide, left, top, card_width, card_height, metric, formatted, delta)

    return slide


def export_enterprise_summary_table(prs, current_data_by_entity, prior_data_by_entity, time_range, unit=None):
    """Export an enterprise summary table showing all divisions side-by-side."""
    slide = _add_content_slide(
        prs,
        "Enterprise Summary \u2014 All Divisions",
        f"{time_range}  |  Generated {datetime.now().strftime('%B %d, %Y')}"
    )

    if unit is None:
        all_vals = []
        for ent_data in current_data_by_entity.values():
            for m in CURRENCY_METRICS:
                all_vals.append(ent_data.get(m, 0))
        unit = determine_display_unit(all_vals)

    # Table dimensions
    rows = len(METRICS) + 1  # header + metrics
    cols = len(DIVISIONS) + 1  # label column + divisions
    tbl_left = Inches(0.4)
    tbl_top = Inches(1.6)
    tbl_width = Inches(12.2)
    tbl_height = Inches(4.8)

    table_shape = slide.shapes.add_table(rows, cols, tbl_left, tbl_top, tbl_width, tbl_height)
    table = table_shape.table

    # Set column widths
    table.columns[0].width = Inches(2.0)
    div_col_width = Inches((12.2 - 2.0) / len(DIVISIONS))
    for i in range(1, cols):
        table.columns[i].width = int(div_col_width)

    # Header row
    header_cell = table.cell(0, 0)
    header_cell.text = "Metric"
    for j, division in enumerate(DIVISIONS, 1):
        cell = table.cell(0, j)
        cell.text = division

    # Style header row
    for j in range(cols):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(12)
            para.font.bold = True
            para.font.color.rgb = WHITE
            para.font.name = "Times New Roman"
            para.alignment = PP_ALIGN.CENTER

    # Data rows
    metrics_list = list(METRICS.keys())
    for i, metric in enumerate(metrics_list, 1):
        # Metric label
        label_cell = table.cell(i, 0)
        label_cell.text = metric
        label_cell.fill.solid()
        label_cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF7, 0xFC)
        for para in label_cell.text_frame.paragraphs:
            para.font.size = Pt(11)
            para.font.bold = True
            para.font.color.rgb = DARK_BLUE
            para.font.name = "Times New Roman"

        # Division values
        for j, division in enumerate(DIVISIONS, 1):
            cell = table.cell(i, j)
            val = current_data_by_entity.get(division, {}).get(metric, 0)
            prior_val = prior_data_by_entity.get(division, {}).get(metric, 0)
            delta = compute_delta(val, prior_val)
            if metric in CURRENCY_METRICS:
                formatted = format_value(val, metric, unit=unit)
            else:
                formatted = format_value(val, metric)

            if delta is None:
                cell.text = f"{formatted}\nn/a"
            else:
                arrow = "▲" if delta >= 0 else "▼"
                cell.text = f"{formatted}\n{arrow} {abs(delta):.1f}%"

            # Alternate row shading
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)

            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(11)
                para.font.name = "Times New Roman"
                para.alignment = PP_ALIGN.CENTER
            # Color the delta line
            if len(cell.text_frame.paragraphs) > 1:
                delta_para = cell.text_frame.paragraphs[1]
                delta_para.font.size = Pt(9)
                if delta is None:
                    delta_para.font.color.rgb = SUBTLE_GRAY
                else:
                    delta_para.font.color.rgb = POS_GREEN if delta >= 0 else RED

    return slide


def export_full_dashboard(
    current_data_by_entity, prior_data_by_entity,
    comparison_data_by_metric, trend_data_by_metric,
    time_range, df, unit=None
):
    """Export the complete dashboard as a multi-slide PowerPoint using BRLS template."""
    prs = create_presentation()

    if unit is None:
        all_vals = []
        for ent_data in current_data_by_entity.values():
            for m in CURRENCY_METRICS:
                all_vals.append(ent_data.get(m, 0))
        unit = determine_display_unit(all_vals)

    # Slide 1: Branded title slide
    slide = _add_title_slide(prs, "Board Performance Dashboard")
    stx = slide.shapes.add_textbox(Inches(5.5), Inches(5.2), Inches(7), Inches(1.5))
    stx.text_frame.word_wrap = True
    p = stx.text_frame.paragraphs[0]
    p.text = f"Time Period: {time_range}"
    p.font.size = Pt(18)
    p.font.color.rgb = DARK_BLUE
    p.font.name = "Times New Roman"
    p.font.bold = True
    p2 = stx.text_frame.add_paragraph()
    p2.text = f"Report Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    p2.font.size = Pt(14)
    p2.font.color.rgb = SUBTLE_GRAY
    p2.font.name = "Times New Roman"
    p3 = stx.text_frame.add_paragraph()
    p3.text = "Divisions: " + "  |  ".join(DIVISIONS)
    p3.font.size = Pt(13)
    p3.font.color.rgb = SUBTLE_GRAY
    p3.font.name = "Times New Roman"

    # Slide 2: Company executive summary KPIs
    if COMPANY_NAME in current_data_by_entity:
        export_executive_summary(
            prs,
            current_data_by_entity[COMPANY_NAME],
            prior_data_by_entity.get(COMPANY_NAME, {}),
            time_range, COMPANY_NAME, unit=unit
        )

    # Slide 3: Enterprise summary table
    export_enterprise_summary_table(prs, current_data_by_entity, prior_data_by_entity, time_range, unit=unit)

    # Division comparison charts for key metrics
    for metric in ["Revenue", "EBITDA", "Pipeline", "Utilization", "Backlog"]:
        if metric in comparison_data_by_metric:
            export_division_comparison(
                prs, comparison_data_by_metric[metric], metric, time_range, unit=unit
            )

    # Trend charts
    for metric in ["Revenue", "Pipeline", "Win Rate", "Utilization", "Backlog"]:
        if metric in trend_data_by_metric:
            export_trend_chart(
                prs, trend_data_by_metric[metric], metric, time_range
            )

    # Individual division detail slides
    for division in DIVISIONS:
        if division in current_data_by_entity:
            export_division_detail(
                prs,
                current_data_by_entity[division],
                prior_data_by_entity.get(division, {}),
                division, time_range, unit=unit
            )

    return prs


def export_single_division(
    division, current_data, prior_data,
    trend_data_by_metric, time_range, unit=None
):
    """Export a deep-dive deck for a single division with KPIs and all trend charts."""
    prs = create_presentation()

    if unit is None:
        currency_vals = [current_data.get(m, 0) for m in CURRENCY_METRICS]
        unit = determine_display_unit(currency_vals)

    # Slide 1: Title slide
    slide = _add_title_slide(prs, f"{division}")
    stx = slide.shapes.add_textbox(Inches(5.5), Inches(5.2), Inches(7), Inches(1.5))
    stx.text_frame.word_wrap = True
    p = stx.text_frame.paragraphs[0]
    p.text = f"Division Performance Report"
    p.font.size = Pt(20)
    p.font.color.rgb = DARK_BLUE
    p.font.name = "Times New Roman"
    p.font.bold = True
    p2 = stx.text_frame.add_paragraph()
    p2.text = f"Time Period: {time_range}  |  Generated {datetime.now().strftime('%B %d, %Y')}"
    p2.font.size = Pt(14)
    p2.font.color.rgb = SUBTLE_GRAY
    p2.font.name = "Times New Roman"

    # Slide 2: Division KPI summary
    export_division_detail(prs, current_data, prior_data, division, time_range, unit=unit)

    # Slides 3+: Trend chart for every metric
    for metric in METRICS:
        trend = trend_data_by_metric.get(metric)
        if trend is not None and not trend.empty:
            slide = _add_content_slide(
                prs,
                f"{division}: {metric} Trend",
                time_range
            )
            fig = _make_trend_chart_fig({division: trend}, metric)
            img_stream = _chart_to_image_bytes(fig, width=1100, height=480)
            if img_stream:
                slide.shapes.add_picture(
                    img_stream, Inches(0.8), Inches(1.6),
                    width=Inches(11.2), height=Inches(5.0)
                )

    return prs


def save_presentation_to_bytes(prs):
    """Save presentation to bytes for download."""
    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output
